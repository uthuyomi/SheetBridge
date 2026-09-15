from pathlib import Path
from typing import Any

from openpyxl import load_workbook

class TemplateAnalyzer:
    
    def analyze_workbook(
        self,
        file_path: str,
    ) -> dict[str, Any]:
        
        path = Path(file_path)
        
        workbook = load_workbook(
            path,
            data_only=False,
        )
        
        sheets = []
        
        for worksheet in workbook.worksheets:
            
            cells = []
            
            for row in worksheet.iter_rows():
                for cell in row:
                    
                    cells.append({
                        "coordinate": cell.coordinate,
                        "value": cell.value,
                        "data_type": cell.data_type,
                        "number_format": cell.number_format, 
                    })
                    
            sheets.append({
                "name": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "merged_cells": [
                    str(cell_range)
                    for cell_range
                    in worksheet.merged_cells.ranges
                ],
                "cells": cells,
            })
            
        return {
            "file_name": path.name,
            "sheets": sheets,
        }
        
    
    def compare_workbooks(
        self,
        template_path: str,
        example_path: str,
    ) -> list[dict[str, Any]]:
        
        template_workbook = load_workbook(
            template_path,
            data_only=False,
        )
        
        example_workbook = load_workbook(
            example_path,
            data_only=False,
        )
        
        differences = []
        
        for template_sheet in template_workbook.worksheets:
            
            if template_sheet.title not in example_workbook.sheetnames:
                continue
            
            example_sheet = example_workbook[
                template_sheet.title
            ]
            
            max_row = max(
                template_sheet.max_row,
                example_sheet.max_row,
            )
            
            max_column = max(
                template_sheet.max_column,
                example_sheet.max_column,
            )
            
            for row in range(1, max_row + 1):
                for column in range(1, max_column +1):
                    
                    template_cell = template_sheet.cell(
                        row=row,
                        column=column,
                    )
                    
                    example_cell = example_sheet.cell(
                        row=row,
                        column=column,
                    )
                    
                    if template_cell.value == example_cell.value:
                        continue
                    
                    differences.append({
                        "sheet": template_sheet.title,
                        "coordinate": template_cell.coordinate,
                        "template_value": template_cell.value,
                        "example_value": example_cell.value,
                        "number_format": example_cell.number_format,
                    })
                    
        return differences